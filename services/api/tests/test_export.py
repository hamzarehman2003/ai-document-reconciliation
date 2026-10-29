"""Exporting approved invoices to Excel and CSV."""

import csv
import io
from datetime import date

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import AuditAction, AuditEvent, AuditObject, Document, MatchCase
from app.services.matching import update_matches
from app.services.review import Decision, decide
from factories import add_three_way_set, unique


def invoice_case(
    session: Session,
    vendor: str,
    decision: Decision | None,
    day: date = date(2026, 3, 2),
) -> MatchCase:
    """A matched invoice (8 billed, 8 received), decided as given."""
    _, _, invoice = add_three_way_set(
        session, ordered="10", received="8", billed="8", vendor=vendor, issue_date=day
    )
    update_matches(session, invoice)
    case = invoice.match_case
    if decision is Decision.APPROVE:
        decide(session, case, decision, "Sara (AP team)", "Checked")
    elif decision is Decision.REJECT:
        decide(session, case, decision, "Sara (AP team)", "Wrong vendor")
    session.flush()
    return case


def excel(client: TestClient, **params: str) -> dict[str, list[tuple]]:
    response = client.get("/api/exports/approved.xlsx", params=params)
    assert response.status_code == 200
    assert response.headers["content-disposition"].startswith(
        'attachment; filename="approved-invoices-'
    )
    workbook = load_workbook(io.BytesIO(response.content))
    return {sheet.title: list(sheet.iter_rows(values_only=True)) for sheet in workbook}


def test_excel_has_approved_invoices_and_their_lines(
    client: TestClient, session: Session
) -> None:
    vendor = unique("Testco")
    approved = invoice_case(session, vendor, Decision.APPROVE)
    invoice_case(session, vendor, None)  # pending: not exported
    invoice_case(session, vendor, Decision.REJECT)  # rejected: not exported

    sheets = excel(client, vendor=vendor)

    header, *rows = sheets["Invoices"]
    assert header[0] == "Invoice number"
    [row] = rows
    invoice: Document = approved.invoice
    assert row[0] == invoice.document_number
    assert row[1] == vendor
    assert row[2].date() == date(2026, 3, 2)  # a real date, not text
    assert row[5] == "EUR"
    assert row[8] == 40  # a real number: 8 × 5.00
    assert row[9] == "Sara (AP team)"
    assert row[11] == "Checked"

    _, *lines = sheets["Line items"]
    assert lines == [(invoice.document_number, "A-1", "Copy paper", 8, 5, 40)]


def test_filters_and_an_empty_export(client: TestClient, session: Session) -> None:
    vendor = unique("Testco")
    invoice_case(session, vendor, Decision.APPROVE, day=date(2026, 3, 2))
    invoice_case(session, vendor, Decision.APPROVE, day=date(2026, 5, 20))

    march = excel(client, vendor=vendor, date_from="2026-03-01", date_to="2026-03-31")
    nothing = excel(client, vendor=unique("Nobody"))

    assert len(march["Invoices"]) == 2  # header + the March invoice
    assert nothing["Invoices"] == [march["Invoices"][0]]  # header only


def test_csv_export(client: TestClient, session: Session) -> None:
    vendor = unique("Testco")
    case = invoice_case(session, vendor, Decision.APPROVE)

    response = client.get("/api/exports/approved.csv", params={"vendor": vendor})

    assert response.content.startswith(b"\xef\xbb\xbf")  # BOM: Excel reads it as UTF-8
    rows = list(csv.DictReader(io.StringIO(response.content.decode("utf-8-sig"))))
    assert len(rows) == 1
    assert rows[0]["Invoice number"] == case.invoice.document_number
    assert rows[0]["Total"] == "40.00"
    assert rows[0]["Invoice date"] == "2026-03-02"


def test_exported_invoices_show_it_in_their_history(
    client: TestClient, session: Session
) -> None:
    vendor = unique("Testco")
    case = invoice_case(session, vendor, Decision.APPROVE)

    client.get(
        "/api/exports/approved.xlsx", params={"vendor": vendor, "reviewer": "Omar"}
    )

    event = session.scalars(
        select(AuditEvent).where(
            AuditEvent.object_type == AuditObject.MATCH_CASE,
            AuditEvent.object_id == case.id,
            AuditEvent.action == AuditAction.CASE_EXPORTED,
        )
    ).one()
    assert event.actor == "Omar"
    assert event.details == {"format": "xlsx"}
