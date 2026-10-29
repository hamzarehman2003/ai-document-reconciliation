"""Export approved invoices as Excel or CSV, ready for the accounting system."""

import csv
import io
from dataclasses import dataclass
from datetime import UTC, date, datetime
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import select
from sqlalchemy.orm import Session, selectinload

from app.models import (
    AuditAction,
    AuditObject,
    Document,
    MatchCase,
    ReviewStatus,
)
from app.services.audit import record

MONEY = "#,##0.00"
QUANTITY = "#,##0.###"
DAY = "yyyy-mm-dd"
TIME = "yyyy-mm-dd hh:mm"

Cell = str | Decimal | date | datetime | None


@dataclass(frozen=True)
class Column:
    title: str
    width: int
    number_format: str | None = None


INVOICE_COLUMNS = [
    Column("Invoice number", 18),
    Column("Vendor", 30),
    Column("Invoice date", 13, DAY),
    Column("Due date", 13, DAY),
    Column("PO number", 16),
    Column("Currency", 10),
    Column("Subtotal", 14, MONEY),
    Column("Tax", 12, MONEY),
    Column("Total", 14, MONEY),
    Column("Approved by", 20),
    Column("Approved at (UTC)", 18, TIME),
    Column("Review note", 40),
    Column("Problems approved anyway", 30),
]
LINE_COLUMNS = [
    Column("Invoice number", 18),
    Column("SKU", 14),
    Column("Description", 40),
    Column("Quantity", 10, QUANTITY),
    Column("Unit price", 12, MONEY),
    Column("Line total", 14, MONEY),
]


def find_approved(
    session: Session,
    date_from: date | None = None,
    date_to: date | None = None,
    vendor: str | None = None,
) -> list[MatchCase]:
    """Approved invoices, oldest invoice date first, with optional filters."""
    query = (
        select(MatchCase)
        .join(MatchCase.invoice)
        .where(MatchCase.review_status == ReviewStatus.APPROVED)
        .options(
            selectinload(MatchCase.invoice).selectinload(Document.line_items),
            selectinload(MatchCase.discrepancies),
        )
        .order_by(Document.issue_date, Document.id)
    )
    if date_from is not None:
        query = query.where(Document.issue_date >= date_from)
    if date_to is not None:
        query = query.where(Document.issue_date <= date_to)
    if vendor:
        query = query.where(Document.vendor_name.icontains(vendor, autoescape=True))
    return list(session.scalars(query))


def invoice_row(case: MatchCase) -> list[Cell]:
    invoice = case.invoice
    problems = sorted({problem.code.value for problem in case.discrepancies})
    return [
        invoice.document_number,
        invoice.vendor_name,
        invoice.issue_date,
        invoice.due_date,
        invoice.po_reference,
        invoice.currency,
        invoice.subtotal,
        invoice.tax,
        invoice.total,
        case.reviewed_by,
        _utc(case.reviewed_at),
        case.review_note,
        ", ".join(problems) or None,
    ]


def line_rows(case: MatchCase) -> list[list[Cell]]:
    return [
        [
            case.invoice.document_number,
            item.sku,
            item.description,
            item.quantity,
            item.unit_price,
            item.line_total,
        ]
        for item in case.invoice.line_items
    ]


def to_xlsx(cases: list[MatchCase]) -> bytes:
    """An Excel file: an "Invoices" sheet and a "Line items" sheet."""
    workbook = Workbook()
    invoices = workbook.active
    invoices.title = "Invoices"
    _fill_sheet(invoices, INVOICE_COLUMNS, [invoice_row(case) for case in cases])
    lines = workbook.create_sheet("Line items")
    _fill_sheet(lines, LINE_COLUMNS, [row for case in cases for row in line_rows(case)])

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def to_csv(cases: list[MatchCase]) -> bytes:
    """The invoices as CSV. UTF-8 with a BOM, so Excel shows € and £ correctly."""
    text = io.StringIO()
    writer = csv.writer(text)
    writer.writerow([column.title for column in INVOICE_COLUMNS])
    for case in cases:
        writer.writerow([_as_text(value) for value in invoice_row(case)])
    return text.getvalue().encode("utf-8-sig")


def record_export(
    session: Session, cases: list[MatchCase], actor: str, file_format: str
) -> None:
    """Each exported invoice's history shows that it left the system."""
    for case in cases:
        record(
            session,
            actor,
            AuditAction.CASE_EXPORTED,
            AuditObject.MATCH_CASE,
            case.id,
            {"format": file_format},
        )


def _fill_sheet(
    sheet: Worksheet, columns: list[Column], rows: list[list[Cell]]
) -> None:
    sheet.append([column.title for column in columns])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append(row)

    for index, column in enumerate(columns, start=1):
        letter = get_column_letter(index)
        sheet.column_dimensions[letter].width = column.width
        if column.number_format:
            for (cell,) in sheet.iter_rows(min_row=2, min_col=index, max_col=index):
                cell.number_format = column.number_format
    sheet.freeze_panes = "A2"  # the header stays visible while scrolling
    sheet.auto_filter.ref = sheet.dimensions  # filter buttons on the header


def _utc(value: datetime | None) -> datetime | None:
    """Excel can't store time zones: write the time in UTC, without the zone."""
    return value.astimezone(UTC).replace(tzinfo=None) if value else None


def _as_text(value: Cell) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(sep=" ", timespec="minutes")
    if isinstance(value, date):
        return value.isoformat()
    return str(value)
